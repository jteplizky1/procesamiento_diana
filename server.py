from __future__ import annotations

import argparse
import hashlib
import io
import itertools
import json
import math
import os
import random
import re
import threading
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import uuid
import webbrowser
from collections import Counter, defaultdict
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

from text_prompts import category_request, classification_request, codebook_key, OPTIONS, PROMPT_VERSION
from brand_rules import brand_memory, normalize_brand, resolve_brand, brand_key
from brand_quality import review_doubtful
import external_review
import cross_analysis
import storage_sync
import gcs_projects
from ai_provider import gemini_request, test_gemini
from text_resilience import run_with_retries, StopRequested, http_status, RETRY_ATTEMPT, OutputLimitError, ModelInterruptedError

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


APP_DIR = Path(__file__).resolve().parent
V1_PROJECTS = APP_DIR.parent / "survey_explorer" / "survey_projects"
PROJECTS_DIR = APP_DIR / "projects"
SNAPSHOTS_DIR = APP_DIR / "snapshots"
JOBS_DIR = APP_DIR / "jobs"
PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
SNAPSHOTS_DIR.mkdir(parents=True, exist_ok=True)
JOBS_DIR.mkdir(parents=True, exist_ok=True)
STATIC_DIR = APP_DIR / "static"
DATA_CACHE: dict[str, pd.DataFrame] = {}
TEXT_JOBS: dict[str, dict] = {}
OLLAMA_ACTIVITY = threading.local()
LOCK = threading.RLock()
CLOUD_SYNC_LOCK = threading.RLock()
CLOUD_SYNC_TIMERS: dict[str, threading.Timer] = {}
CLOUD_SYNC_STATUS: dict[str, dict] = {}
CLOUD_HYDRATION = {"attempted_at": 0.0, "complete": False, "error": ""}
JOB_SYNC_TIMERS: dict[str, threading.Timer] = {}

QUESTION_TYPES = [
    "selección única", "selección múltiple", "escala", "matriz de escala",
    "ordenamiento/ranking", "texto libre", "demográfica", "numérica",
    "identificador", "técnica", "vacía",
]


def clean_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): clean_json(v) for k, v in value.items()}
    if isinstance(value, list):
        return [clean_json(v) for v in value]
    if pd.isna(value) if not isinstance(value, (list, dict)) else False:
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def project_path(project_id: str) -> Path:
    return PROJECTS_DIR / f"{project_id}.json"


def job_path(job_id: str) -> Path:
    return JOBS_DIR / f"{job_id}.json"


def _write_job(job: dict) -> None:
    job["updated_at"] = time.time()
    destination = job_path(job["id"])
    temporary = destination.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(clean_json(job), ensure_ascii=False), encoding="utf-8")
    os.replace(temporary, destination)


def _upload_job(job_id: str) -> None:
    try:
        path = job_path(job_id)
        if path.exists():
            gcs_projects.save_job(json.loads(path.read_text(encoding="utf-8")))
    finally:
        with CLOUD_SYNC_LOCK:
            JOB_SYNC_TIMERS.pop(job_id, None)


def persist_job(job: dict, *, immediate_cloud: bool = False) -> None:
    """Store progress locally and in GCS so another instance can answer polling."""
    _write_job(job)
    if immediate_cloud:
        gcs_projects.save_job(clean_json(job))
        return
    with CLOUD_SYNC_LOCK:
        previous = JOB_SYNC_TIMERS.pop(job["id"], None)
        if previous:
            previous.cancel()
        timer = threading.Timer(.35, _upload_job, args=(job["id"],))
        timer.daemon = True
        JOB_SYNC_TIMERS[job["id"]] = timer
        timer.start()


def read_job(job_id: str) -> dict | None:
    if job_id in TEXT_JOBS:
        return TEXT_JOBS[job_id]
    path = job_path(job_id)
    if path.exists():
        local = json.loads(path.read_text(encoding="utf-8"))
        if local.get("status") != "running":
            return local
        # This replica is only observing the worker. Refresh its moving state from GCS.
        try:
            remote = gcs_projects.load_job(job_id)
        except Exception:
            return local
        if remote:
            path.write_text(json.dumps(remote, ensure_ascii=False), encoding="utf-8")
            return remote
        return local
    job = gcs_projects.load_job(job_id)
    if job:
        path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    return job


def delete_project(project_id: str) -> dict:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", project_id):
        raise ValueError("ID de proyecto inválido")
    with LOCK:
        if any(job.get("project_id") == project_id and job.get("status") == "running" for job in TEXT_JOBS.values()):
            raise ValueError("No se puede eliminar el proyecto mientras procesa textos. Detené la cola primero.")
        targets = [project_path(project_id), SNAPSHOTS_DIR / f"{project_id}.jsonl",
                   V1_PROJECTS / f"{project_id}.json"]
        removed = []
        for path in targets:
            if path.exists():
                path.unlink()
                removed.append(path.name)
        DATA_CACHE.pop(project_id, None)
        for job_id in [job_id for job_id, job in TEXT_JOBS.items() if job.get("project_id") == project_id]:
            TEXT_JOBS.pop(job_id, None)
    if not removed:
        raise FileNotFoundError("Proyecto inexistente")
    return {"deleted": True, "removed_files": len(removed)}


def read_project(project_id: str) -> dict:
    path = project_path(project_id)
    if not path.exists():
        legacy = V1_PROJECTS / f"{project_id}.json"
        if legacy.exists():
            data = json.loads(legacy.read_text(encoding="utf-8"))
            data.setdefault("version", 2)
            save_project(data)
            return data
        raise FileNotFoundError("Proyecto inexistente")
    return json.loads(path.read_text(encoding="utf-8"))


def _upload_project_checkpoint(project_id: str) -> None:
    """Upload the latest durable project state without blocking an HTTP request."""
    try:
        latest = read_project(project_id)
        result = gcs_projects.save_bundle(
            latest, project_path(project_id), SNAPSHOTS_DIR / f"{project_id}.jsonl", None
        )
        CLOUD_SYNC_STATUS[project_id] = {
            "ok": True, "saved_at": datetime.now().isoformat(timespec="seconds"), **result
        }
    except Exception as error:
        CLOUD_SYNC_STATUS[project_id] = {
            "ok": False, "error": str(error),
            "attempted_at": datetime.now().isoformat(timespec="seconds"),
        }
    finally:
        with CLOUD_SYNC_LOCK:
            CLOUD_SYNC_TIMERS.pop(project_id, None)


def schedule_cloud_checkpoint(project_id: str, delay: float = 1.0) -> None:
    """Debounce rapid edits and persist the newest version using the app identity."""
    if os.getenv("DIANA_AUTO_CLOUD", "1").strip().lower() in {"0", "false", "no"}:
        return
    with CLOUD_SYNC_LOCK:
        previous = CLOUD_SYNC_TIMERS.pop(project_id, None)
        if previous:
            previous.cancel()
        timer = threading.Timer(delay, _upload_project_checkpoint, args=(project_id,))
        timer.daemon = True
        CLOUD_SYNC_TIMERS[project_id] = timer
        timer.start()


def hydrate_cloud_projects(force: bool = False) -> dict:
    """Recover cloud projects after a cold start of an ephemeral server."""
    now = time.time()
    if not force and (CLOUD_HYDRATION["complete"] or now - CLOUD_HYDRATION["attempted_at"] < 60):
        return dict(CLOUD_HYDRATION)
    CLOUD_HYDRATION["attempted_at"] = now
    try:
        restored = 0
        for remote in gcs_projects.list_cloud_projects():
            if not project_path(remote["id"]).exists():
                gcs_projects.restore(remote["prefix"], PROJECTS_DIR, SNAPSHOTS_DIR)
                restored += 1
            CLOUD_SYNC_STATUS.setdefault(remote["id"], {
                "ok": True, "saved_at": remote.get("updated_at", ""),
                "bucket": f"gs://{gcs_projects.bucket_name()}", "prefix": remote["prefix"],
                "restored": True,
            })
        CLOUD_HYDRATION.update({"complete": True, "error": "", "restored": restored})
    except Exception as error:
        CLOUD_HYDRATION.update({"complete": False, "error": str(error), "restored": 0})
    return dict(CLOUD_HYDRATION)


def save_project(project: dict, *, cloud: bool = True) -> None:
    project["updated_at"] = datetime.now().isoformat(timespec="seconds")
    project_path(project["id"]).write_text(
        json.dumps(clean_json(project), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if cloud:
        schedule_cloud_checkpoint(project["id"])


def list_projects() -> list[dict]:
    found: dict[str, dict] = {}
    for folder in (V1_PROJECTS, PROJECTS_DIR):
        if not folder.exists():
            continue
        for path in folder.glob("*.json"):
            try:
                item = json.loads(path.read_text(encoding="utf-8"))
                found[item["id"]] = {
                    "id": item["id"], "name": item.get("name", "Sin nombre"),
                    "description": item.get("description", ""),
                    "updated_at": item.get("updated_at", item.get("created_at", "")),
                    "source_url": item.get("source_url", ""),
                }
            except Exception:
                continue
    return sorted(found.values(), key=lambda x: x.get("updated_at", ""), reverse=True)


def spreadsheet_id(url: str) -> str:
    match = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", url)
    if not match:
        raise ValueError("El enlace no parece corresponder a Google Sheets")
    return match.group(1)


def google_csv_url(url: str, selected_gid: str | int | None = None) -> str:
    gid = str(selected_gid) if selected_gid is not None and str(selected_gid) != "" else None
    if gid is None:
        match = re.search(r"[?#&]gid=(\d+)", url)
        gid = match.group(1) if match else None
    suffix = f"&gid={gid}" if gid is not None else ""
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id(url)}/export?format=csv{suffix}"


def list_google_sheet_tabs(url: str) -> list[dict]:
    """Obtiene nombres y gid desde el bootstrap público de Google Sheets."""
    edit_url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id(url)}/edit"
    request = urllib.request.Request(edit_url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=30) as response:
        html = response.read().decode("utf-8", errors="replace")
    marker = "var bootstrapData = "
    start = html.find(marker)
    if start < 0:
        raise ValueError("No fue posible leer las solapas. Verificá que la hoja sea pública.")
    bootstrap, _ = json.JSONDecoder().raw_decode(html[start + len(marker):])
    snapshots = bootstrap.get("changes", {}).get("topsnapshot", [])
    tabs = []
    for entry in snapshots:
        if not isinstance(entry, list) or len(entry) < 2 or entry[0] != 21350203:
            continue
        try:
            payload = json.loads(entry[1])
            gid = str(payload[2])
            title = str(payload[3][0]["1"][0][2])
            tabs.append({"gid": gid, "name": title})
        except (ValueError, TypeError, KeyError, IndexError):
            continue
    if not tabs:
        # En algunas variantes Google sólo renderiza la solapa activa.
        caption = re.search(r'docs-sheet-tab-caption[^>]*>([^<]+)<', html)
        grid = bootstrap.get("gridId")
        if caption and grid is not None:
            tabs.append({"gid": str(grid), "name": caption.group(1).strip()})
    if not tabs:
        raise ValueError("Google no informó ninguna solapa accesible.")
    unique = {item["gid"]: item for item in tabs}
    return list(unique.values())


def normalize_columns(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    seen: Counter = Counter()
    columns = []
    for raw in frame.columns:
        base = str(raw).strip() or "Sin título"
        seen[base] += 1
        columns.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
    frame.columns = columns
    return frame


def resolve_column(frame: pd.DataFrame, requested: str) -> str:
    """Resuelve encabezados que el navegador colapsó por espacios HTML consecutivos."""
    if requested in frame.columns:
        return requested
    canonical = lambda value: re.sub(r"\s+", " ", str(value)).strip()
    matches = [column for column in frame.columns if canonical(column) == canonical(requested)]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError("Hay más de una pregunta con el mismo texto al ignorar espacios. Renombrá los encabezados para diferenciarlas.")
    raise ValueError("La pregunta no existe en la muestra. Actualizá la fuente y revisá el diccionario de preguntas.")


def read_google_csv(url: str) -> pd.DataFrame:
    """Descarga una solapa pública y convierte errores de Google en mensajes útiles."""
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 SurveyStudio/2.0", "Accept": "text/csv,*/*"},
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            payload = response.read()
            content_type = response.headers.get("Content-Type", "").lower()
    except urllib.error.HTTPError as error:
        if error.code in (401, 403):
            raise ValueError(
                "Google no permite descargar esta hoja sin iniciar sesión. "
                "Abrí la hoja en Google Sheets y elegí Compartir > Acceso general > "
                "Cualquier persona con el enlace > Lector. Después volvé a actualizar la fuente. "
                "Si esa opción no aparece, la política de Google Workspace de tu empresa bloquea "
                "el acceso público; en ese caso usá una copia autorizada para compartir."
            ) from None
        if error.code == 404:
            raise ValueError(
                "Google no encontró la hoja. Revisá que el enlace sea correcto y que el archivo no haya sido eliminado."
            ) from None
        raise ValueError(f"Google devolvió un error HTTP {error.code} al descargar la hoja.") from None
    except urllib.error.URLError as error:
        reason = getattr(error, "reason", error)
        raise ValueError(f"No se pudo conectar con Google Sheets: {reason}") from None
    except TimeoutError:
        raise ValueError("Google Sheets tardó demasiado en responder. Intentá actualizar nuevamente.") from None

    # Una pantalla de inicio de sesión puede llegar con estado 200 en algunos flujos.
    preview = payload[:2000].lower()
    if "text/html" in content_type and (b"accounts.google" in preview or b"signin" in preview):
        raise ValueError(
            "Google mostró una pantalla de inicio de sesión en vez de los datos. "
            "Configurá Compartir > Acceso general > Cualquier persona con el enlace > Lector."
        )
    try:
        return pd.read_csv(io.BytesIO(payload))
    except (pd.errors.EmptyDataError, UnicodeDecodeError, pd.errors.ParserError) as error:
        raise ValueError(f"Google respondió, pero la solapa no contiene un CSV válido: {error}") from None


def load_source(project: dict, refresh: bool = False) -> pd.DataFrame:
    pid = project["id"]
    snapshot = SNAPSHOTS_DIR / f"{pid}.jsonl"
    if not refresh and pid in DATA_CACHE:
        return DATA_CACHE[pid]
    if not refresh and snapshot.exists():
        frame = pd.read_json(snapshot, lines=True, dtype=False)
        DATA_CACHE[pid] = frame
        return frame
    url = project.get("source_url", "").strip()
    if not url:
        raise ValueError("El proyecto todavía no tiene una fuente")
    frame = normalize_columns(read_google_csv(google_csv_url(url, project.get("source_sheet_gid"))))
    frame.to_json(snapshot, orient="records", lines=True, force_ascii=False)
    DATA_CACHE[pid] = frame
    project["source_rows"] = len(frame)
    project["source_columns"] = list(frame.columns)
    project["source_refreshed_at"] = datetime.now().isoformat(timespec="seconds")
    project["dictionary"] = merge_dictionary(project.get("dictionary", []), build_dictionary(frame))
    save_project(project)
    return frame


def norm(value: Any) -> str:
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFKD", str(value).lower().strip()) if not unicodedata.combining(c))


def split_answers(value: Any) -> list[str]:
    if value is None or pd.isna(value):
        return []
    text = str(value).strip()
    if not text:
        return []
    return [x.strip() for x in re.split(r"\r?\n+|\s*;\s*", text) if x.strip()]


def parse_ranking(value: Any) -> list[tuple[str, int]]:
    result = []
    for part in split_answers(value):
        patterns = [r"^\s*(\d+)\s*[.):-]\s*(.+?)\s*$", r"^\s*(.+?)\s*[:=-]\s*(\d+)\s*$"]
        match = re.match(patterns[0], part)
        if match:
            result.append((match.group(2).strip(), int(match.group(1))))
            continue
        match = re.match(patterns[1], part)
        if match:
            result.append((match.group(1).strip(), int(match.group(2))))
    return result


def parse_matrix_scale(value: Any, minimum: int = 1, maximum: int = 5) -> list[tuple[str, int]]:
    result = []
    for part in split_answers(value):
        match = re.match(r"^\s*(.+?)\s*:\s*(\d+)\s*$", part)
        if not match:
            return []
        score = int(match.group(2))
        if not minimum <= score <= maximum:
            return []
        result.append((match.group(1).strip(), score))
    return result if len(result) >= 2 else []


def detect_type(series: pd.Series, name: str) -> str:
    clean = series.dropna()
    n = norm(name)
    if any(x in n for x in ["email", "correo", "nombre", "apellido", "submission id", "user id"]):
        return "identificador"
    if any(x in n for x in ["timestamp", "fecha de envio", "referer", "url track"]):
        return "técnica"
    if any(x in n for x in ["genero", "sexo", "edad", "region", "ingreso", "comuna", "ciudad", "ocupacion"]):
        return "demográfica"
    if clean.empty:
        return "vacía"
    if pd.api.types.is_numeric_dtype(clean):
        return "escala" if clean.nunique() <= 11 else "numérica"
    sample = clean.astype(str).head(200)
    matrix_rows = [parse_matrix_scale(x) for x in sample]
    structured = [row for row in matrix_rows if row]
    if len(structured) / max(len(sample), 1) >= .2:
        return "matriz de escala"
    if sample.map(lambda x: len(parse_ranking(x)) >= 2).mean() >= .2:
        return "ordenamiento/ranking"
    multiline = sample.str.contains(r"\r?\n", regex=True).mean()
    multiline_parts = [split_answers(x) for x in sample if len(split_answers(x)) >= 2]
    short_options = multiline_parts and sum(len(p) for row in multiline_parts for p in row) / sum(len(row) for row in multiline_parts) <= 100
    unique = clean.nunique() / max(len(clean), 1)
    if multiline >= .2 and short_options:
        return "selección múltiple"
    if unique >= .45 or clean.nunique() > 40:
        return "texto libre"
    return "selección única"


def build_dictionary(frame: pd.DataFrame) -> list[dict]:
    result = []
    for col in frame.columns:
        detected = detect_type(frame[col], col)
        result.append({"pregunta": col, "tipo": detected, "tipo_detectado": detected,
                       "respuestas": int(frame[col].notna().sum()),
                       "valores_unicos": int(frame[col].nunique(dropna=True)),
                       "incluir": detected not in {"identificador", "técnica", "vacía"}})
    return result


def merge_dictionary(saved: list[dict], detected: list[dict]) -> list[dict]:
    old = {x.get("pregunta"): x for x in saved}
    return [{**row, **{k: v for k, v in old.get(row["pregunta"], {}).items() if k in {"tipo", "incluir"}}} for row in detected]


def response_ids(frame: pd.DataFrame) -> list[str]:
    candidates = [c for c in frame.columns if any(x in norm(c) for x in ["submission id", "response id", "marca temporal", "timestamp"])]
    if candidates and frame[candidates[0]].notna().all() and frame[candidates[0]].astype(str).is_unique:
        return frame[candidates[0]].astype(str).tolist()
    return [hashlib.sha256(json.dumps(clean_json(row), ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:20] for row in frame.to_dict("records")]


def quota_status(frame: pd.DataFrame, dimensions: list[str], rows: list[dict], mode: str, total: int) -> dict:
    if not dimensions:
        raise ValueError("Elegí al menos una variable de cuota")
    value_key = "cantidad" if mode == "Cantidades" else "porcentaje"
    values = [float(row.get(value_key, 0) or 0) for row in rows]
    expected = total if mode == "Cantidades" else 100
    if abs(sum(values) - expected) > .01:
        raise ValueError(f"La suma es {sum(values):g}; debe ser {expected:g}")
    exact = values if mode == "Cantidades" else [x * total / 100 for x in values]
    targets = [int(math.floor(x)) for x in exact]
    for idx in sorted(range(len(exact)), key=lambda i: exact[i] - targets[i], reverse=True)[: total - sum(targets)]:
        targets[idx] += 1
    output = []
    for row, target in zip(rows, targets):
        mask = pd.Series(True, index=frame.index)
        for d in dimensions:
            mask &= frame[d].fillna("").astype(str).eq(str(row.get(d, "")))
        available = int(mask.sum())
        output.append({**{d: row.get(d, "") for d in dimensions}, "objetivo": target, "disponibles": available,
                       "faltantes": max(target - available, 0), "excedentes": max(available - target, 0),
                       "cumplimiento": round(min(available, target) / total * 100, 2) if total else 0})
    return {"rows": output, "total": total, "available": sum(min(x["objetivo"], x["disponibles"]) for x in output),
            "missing": sum(x["faltantes"] for x in output), "discard": sum(x["excedentes"] for x in output)}


def make_quota_rows(frame: pd.DataFrame, dimensions: list[str], total: int, mode: str) -> list[dict]:
    if not dimensions:
        raise ValueError("Elegí al menos una variable de cuota")
    missing = [dimension for dimension in dimensions if dimension not in frame]
    if missing:
        raise ValueError("No existen estas variables en la fuente: " + ", ".join(missing))
    values = []
    for dimension in dimensions:
        options = [clean_json(value) for value in frame[dimension].dropna().unique().tolist()
                   if str(value).strip()]
        options.sort(key=lambda value: norm(value))
        if not options:
            raise ValueError(f"La variable {dimension} no tiene valores para combinar")
        values.append(options)
    key = "porcentaje" if mode == "Porcentajes" else "cantidad"
    return [{**dict(zip(dimensions, combination)), key: 0}
            for combination in itertools.product(*values)]


def balanced_sample(frame: pd.DataFrame, project: dict, seed: int, weighted: bool = False) -> tuple[list[str], int, dict[str, float]]:
    settings = project.get("quota_settings", {})
    dims = settings.get("dimensions", [])
    status = quota_status(frame, dims, project.get("quota_rows", []), settings.get("input_mode", "Cantidades"), int(settings.get("target_total", 0)))
    zero = [q for q in status['rows'] if q['objetivo'] > 0 and q['disponibles'] == 0]
    if zero and weighted:
        labels = [' / '.join(str(q.get(d, '')) for d in dims) for q in zero]
        raise ValueError('No se puede ponderar una cuota objetivo sin respuestas: ' + '; '.join(labels))
    if status["missing"] and not weighted:
        raise ValueError("No hay respuestas suficientes para completar todas las cuotas")
    ids = response_ids(frame)
    selected = []
    weights = {}
    rng = random.Random(seed)
    for q in status["rows"]:
        candidates = []
        for idx, row in frame.iterrows():
            if all(str(row.get(d, "")) == str(q.get(d, "")) for d in dims):
                candidates.append(idx)
        take = min(len(candidates), q['objetivo'])
        picked = rng.sample(candidates, take)
        selected.extend(picked)
        weight = q['objetivo'] / take if take else 0
        for idx in picked:
            weights[ids[idx]] = round(weight, 8)
    return [ids[i] for i in selected], len(selected), weights


def expanded_sample_ids(frame: pd.DataFrame, project: dict, selected_ids: list[str], seed: int) -> list[str]:
    """Crea una base de tamaño objetivo; sólo replica dentro de una cuota subcumplida."""
    settings = project.get("quota_settings", {})
    dims = settings.get("dimensions", [])
    status = quota_status(frame, dims, project.get("quota_rows", []), settings.get("input_mode", "Cantidades"),
                          int(settings.get("target_total", 0)))
    ids = response_ids(frame)
    selected = set(selected_ids)
    rng = random.Random(seed)
    expanded = []
    for q in status["rows"]:
        available = [rid for rid, (_, row) in zip(ids, frame.iterrows())
                     if rid in selected and all(str(row.get(d, "")) == str(q.get(d, "")) for d in dims)]
        target = q["objetivo"]
        if target and not available:
            labels = " / ".join(str(q.get(d, "")) for d in dims)
            raise ValueError(f"No se puede completar la cuota sin ninguna respuesta: {labels}")
        expanded.extend(available[:target])
        if len(available) < target:
            expanded.extend(rng.choices(available, k=target - len(available)))
    rng.shuffle(expanded)
    return expanded


def sample_frame(frame: pd.DataFrame, project: dict) -> pd.DataFrame:
    chosen = set(project.get("sample_response_ids", []))
    if not chosen:
        return frame.iloc[0:0]
    ids = response_ids(frame)
    return frame.loc[[i for i, rid in enumerate(ids) if rid in chosen]].copy()


def text_workflow_status(frame: pd.DataFrame, project: dict) -> dict:
    sample = sample_frame(frame, project)
    all_ids = response_ids(frame)
    chosen = set(project.get("sample_response_ids", []))
    ids = [rid for rid in all_ids if rid in chosen]
    available = [x["pregunta"] for x in project.get("dictionary", [])
                 if x.get("incluir") and x.get("tipo") == "texto libre" and x["pregunta"] in sample]
    selected = project.get("selected_text_questions")
    questions = [question for question in (selected if selected is not None else available)
                 if question in available]
    details = []
    for question in questions:
        mapping = project.get("text_classifications", {}).get(question, {})
        overrides = project.get("text_response_overrides", {}).get(question, {})
        total = processed = approved = 0
        for rid, raw in zip(ids, sample[question]):
            original = "" if raw is None or bool(pd.isna(raw)) else str(raw)
            if not original.strip():
                continue
            total += 1
            item = overrides.get(rid, mapping.get(original, {}))
            segment = item.get("segment", "")
            if str(segment).strip():
                processed += 1
                if item.get("approved"):
                    approved += 1
        details.append({"pregunta": question, "total": total, "procesadas": processed,
                        "aprobadas": approved, "pendientes": total - processed,
                        "pendientes_revision": total - approved, "procesamiento_completo": processed == total,
                        "completa": approved == total})
    return {"complete": all(x["completa"] for x in details),
            "processing_complete": all(x["procesamiento_completo"] for x in details),
            "questions": details, "total": sum(x["total"] for x in details),
            "processed": sum(x["procesadas"] for x in details),
            "approved": sum(x["aprobadas"] for x in details)}


def reconcile_changed_questions(project: dict, new_dictionary: list[dict]) -> list[str]:
    """Cambia el enrutamiento sin borrar resultados anteriores de la pregunta."""
    old = {x.get("pregunta"): x for x in project.get("dictionary", [])}
    new = {x.get("pregunta"): x for x in new_dictionary}
    changed = [q for q in sorted(set(old) | set(new))
               if (old.get(q, {}).get("tipo"), old.get(q, {}).get("incluir")) !=
                  (new.get(q, {}).get("tipo"), new.get(q, {}).get("incluir"))]
    project["selected_text_questions"] = [q for q in project.get("selected_text_questions", [])
                                          if q not in changed and new.get(q, {}).get("tipo") == "texto libre"]
    project["selected_ranking_questions"] = [q for q in project.get("selected_ranking_questions", [])
                                             if q not in changed and new.get(q, {}).get("tipo") == "ordenamiento/ranking"]
    return changed


def ai_request(project: dict, messages: list[dict], schema: dict) -> dict:
    return gemini_request(project, messages, schema, getattr(OLLAMA_ACTIVITY, "callback", lambda **kw: None))


def process_text(
    project: dict,
    frame: pd.DataFrame,
    question: str,
    batch_size: int,
    mode: str = "semantic",
    category_count: int | None = None,
    progress=None,
) -> dict:
    progress = progress or (lambda percent, phase, **extra: None)
    progress(3, "Preparando respuestas")
    values = sample_frame(frame, project)[question].dropna().astype(str)
    values = values[values.str.strip().ne("")]
    classifications = project.setdefault("text_classifications", {}).setdefault(question, {})
    project.setdefault("text_processing_modes", {})[question] = mode
    project.setdefault("text_category_counts", {})[question] = category_count
    pending = [x for x in values.unique().tolist() if x not in classifications]
    batch = pending[:batch_size]
    total_unique = len(values.unique())
    already_done = total_unique - len(pending)
    if not batch:
        progress(100, "Todo procesado", completed=total_unique, total=total_unique, remaining=0)
        return {"processed": 0, "remaining": 0, "completed": total_unique, "total": total_unique}
    codebooks = project.setdefault("text_codebooks", {})
    instructions = project.get("text_question_settings", {}).get(question, {}).get("instructions", "")
    known, examples = brand_memory(project) if mode == "brands" else ([], {})
    key = codebook_key(question, category_count)
    if mode != "brands" and key not in codebooks:
        request = category_request(question, values.tolist(), category_count, instructions, RETRY_ATTEMPT.get())
        OLLAMA_ACTIVITY.stage = 'Creación de categorías'
        progress(0, f"Creando categorías (muestra de {request['sample_count']} textos completos)", completed=already_done, total=total_unique, remaining=len(pending))
        result = ai_request(project, request["messages"], request["schema"])
        segments = result.get("segments", [])
        if not segments or any(not isinstance(x, dict) or not isinstance(x.get("name"), str) or not x["name"].strip() for x in segments):
            raise ValueError("Gemini devolvió un catálogo inválido.")
        if category_count and len(segments) != category_count:
            raise ValueError(f"Gemini devolvió {len(segments)} categorías; se esperaban {category_count}.")
        codebooks[key] = segments
        with LOCK:
            latest = read_project(project["id"])
            latest.setdefault("text_codebooks", {})[key] = segments
            save_project(latest)
    names = [x["name"] for x in codebooks.get(key, [])]
    request = classification_request(question, batch, mode, names, instructions, known, examples)
    OLLAMA_ACTIVITY.stage = 'Clasificación de respuestas'
    progress(0, f"Clasificando {len(batch)} textos con Gemini", completed=already_done, total=total_unique, remaining=len(pending))
    result = ai_request(project, request["messages"], request["schema"])
    progress(0, "Validando resultados", completed=already_done, total=total_unique, remaining=len(pending))
    valid = {}
    raw_predictions = {}
    for item in result.get("items", []):
        if not isinstance(item, dict):
            continue
        index = item.get("id")
        segment = item.get("segment")
        if isinstance(index, str) and index.isdigit():
            index = int(index)
        if type(index) is not int or not 0 <= index < len(batch):
            continue
        if not isinstance(segment, str) or not segment.strip():
            continue
        if mode == "brands":
            raw_predictions[batch[index]] = segment
            segment = resolve_brand(batch[index], segment, known, examples)
        valid[batch[index]] = {"segment": segment.strip(), "reason": item.get("reason", ""), "approved": False}
    if not valid:
        raise ValueError("Gemini no devolvió clasificaciones válidas. No se guardó un lote vacío; podés reintentar.")
    if mode == "brands":
        valid = review_doubtful(question, valid, raw_predictions, known, examples, instructions,
                               lambda messages, schema: ai_request(project, messages, schema), progress)
    # Merge against the latest project so a concurrent manual review is preserved.
    with LOCK:
        latest = read_project(project["id"])
        saved = latest.setdefault("text_classifications", {}).setdefault(question, {})
        batch_number = max((v.get("batch", 0) for v in saved.values()), default=0) + 1
        for answer in batch:
            if answer in valid and answer not in saved:
                saved[answer] = {**valid[answer], "batch": batch_number}
        if mode != "brands":
            latest.setdefault("text_codebooks", {})[key] = codebooks[key]
        latest.setdefault("text_processing_modes", {})[question] = mode
        latest.setdefault("text_category_counts", {})[question] = category_count
        save_project(latest)
    remaining = sum(answer not in saved for answer in values.unique())
    processed = len(pending) - remaining
    progress(100, "Lote guardado", completed=total_unique - remaining, total=total_unique, remaining=remaining)
    return {"processed": processed, "remaining": remaining, "completed": total_unique - remaining, "total": total_unique, "batch": batch_number}


def text_selection(project, data):
    eligible = {x["pregunta"] for x in project.get("dictionary", []) if x.get("incluir") and x.get("tipo") == "texto libre"}
    questions = data.get("questions") or ([data["question"]] if data.get("question") else [])
    if not questions or len(questions) != len(set(questions)) or any(q not in eligible for q in questions):
        raise ValueError("Seleccioná preguntas de texto libre válidas, sin duplicados.")
    settings = {}
    for question in questions:
        cfg = data.get("settings", {}).get(question, data)
        mode = cfg.get("mode", "semantic")
        count = cfg.get("category_count")
        count = int(count) if count not in (None, "", 0, "0") else None
        if mode not in ("semantic", "brands") or (count is not None and count < 2):
            raise ValueError("Revisá el modo y la cantidad de categorías.")
        settings[question] = {"mode": mode, "category_count": count, "instructions": str(cfg.get("instructions", "")).strip()}
    return questions, settings


def text_prompt_preview(project, frame, data):
    questions, settings = text_selection(project, data)
    question = questions[0]
    cfg = settings[question]
    values = sample_frame(frame, project)[question].dropna().astype(str)
    values = values[values.str.strip().ne("")]
    mapping = project.get("text_classifications", {}).get(question, {})
    batch = [v for v in values.unique() if v not in mapping][:int(data.get("batch_size", 10))]
    if not batch:
        return {"version": PROMPT_VERSION, "requests": [], "note": "No quedan textos pendientes."}
    key = codebook_key(question, cfg["category_count"])
    categories = project.get("text_codebooks", {}).get(key)
    if cfg["mode"] != "brands" and not categories:
        requests = [category_request(question, values.tolist(), cfg["category_count"], cfg["instructions"])]
        note = "Esta es la próxima llamada exacta. El prompt de clasificación se construirá con el catálogo devuelto."
    else:
        known, examples = brand_memory(project) if cfg["mode"] == "brands" else ([], {})
        requests = [classification_request(question, batch, cfg["mode"], [x["name"] for x in (categories or [])], cfg["instructions"], known, examples)]
        note = "Próxima llamada exacta para los textos pendientes."
    return {"version": PROMPT_VERSION, "options": OPTIONS, "requests": requests, "note": note}


def start_text_job(project_id: str, data: dict) -> dict:
    project = read_project(project_id)
    questions, settings = text_selection(project, data)
    batch_size = int(data.get("batch_size", 10))
    if not 1 <= batch_size <= 100:
        raise ValueError("El lote debe tener entre 1 y 100 textos.")
    with LOCK:
        if any(j["status"] == "running" for j in TEXT_JOBS.values()):
            raise ValueError("Ya hay una cola ejecutándose. Esperá o detenela antes de iniciar otra.")
        job_id = uuid.uuid4().hex
        job = {"id": job_id, "project_id": project_id, "status": "running", "percent": 0,
               "phase": "Preparando cola", "questions": questions, "started_at": time.time(),
               "completed": 0, "total": 0, "remaining": 0, "saved_batches": 0, "batches": [], "stop_requested": False}
        TEXT_JOBS[job_id] = job
        project["last_text_job_id"] = job_id
        save_project(project)
        try:
            persist_job(job, immediate_cloud=True)
        except Exception as error:
            job["persistence_warning"] = str(error)
            persist_job(job)

    def update(_percent, phase, **extra):
        with LOCK:
            job["phase"] = phase
            if "completed" in extra:
                job["question_completed"] = extra["completed"]
                job["question_total"] = extra["total"]
            persist_job(job)

    def activity(**extra):
        with LOCK:
            job.update(extra)
            persist_job(job)

    def worker():
        OLLAMA_ACTIVITY.callback = activity
        try:
            frame = load_source(read_project(project_id))
            base = sample_frame(frame, read_project(project_id))
            unique = {q: list(dict.fromkeys(v for v in base[q].dropna().astype(str) if v.strip())) for q in questions}
            def recount():
                current = read_project(project_id)
                total = sum(len(v) for v in unique.values())
                complete = sum(sum(v in current.get("text_classifications", {}).get(q, {}) for v in values) for q, values in unique.items())
                activity(total=total, completed=complete, remaining=total-complete,
                         percent=round(100*complete/total, 1) if total else 100)
            recount()
            failed_questions = []
            for index, question in enumerate(questions, 1):
                activity(question=question, question_index=index, question_count=len(questions))
                while True:
                    if job["stop_requested"]:
                        activity(status="stopped", phase="Cola detenida; resultados guardados")
                        return
                    try:
                        result = run_with_retries(
                            lambda size: process_text(read_project(project_id), frame, question, size,
                                settings[question]["mode"], settings[question]["category_count"], update),
                            batch_size, activity, lambda: job["stop_requested"])
                    except StopRequested:
                        activity(status="stopped", phase="Cola detenida; resultados guardados")
                        return
                    except Exception as error:
                        if http_status(error) in (401, 403):
                            raise
                        failed_questions.append({"question": question, "error": str(error), "details": traceback.format_exc()})
                        activity(failed_questions=list(failed_questions), phase="Pregunta pendiente por error; continúa la siguiente")
                        break
                    if result["processed"]:
                        with LOCK:
                            job["saved_batches"] += 1
                            job["batches"].append({"number": job["saved_batches"], "question": question,
                                                  "processed": result["processed"], "finished_at": time.time()})
                            persist_job(job)
                    recount()
                    if not data.get("all_batches", True) or result["remaining"] == 0:
                        break
                    if result["processed"] == 0:
                        raise ValueError("El lote no produjo nuevos resultados; se detuvo la cola.")
            recount()
            activity(status="partial" if failed_questions else "complete",
                     phase="Finalizó con preguntas pendientes por error" if failed_questions else "Cola finalizada")
        except Exception as error:
            activity(status="error", phase="Cola detenida por error", error=str(error), error_details=traceback.format_exc())
        finally:
            OLLAMA_ACTIVITY.callback = lambda **kw: None
            persist_job(job)

    threading.Thread(target=worker, daemon=True, name=f"text-job-{job_id[:8]}").start()
    return dict(job)


def ranking_table(frame: pd.DataFrame, question: str, weights=None) -> list[dict]:
    weights = pd.Series(weights if weights is not None else 1.0, index=frame.index, dtype=float)
    counts: dict[str, Counter] = defaultdict(Counter)
    for idx, value in frame[question].dropna().items():
        for option, position in parse_ranking(value):
            counts[option][position] += weights.loc[idx]
    result = []
    for option, positions in counts.items():
        base = sum(positions.values())
        row = {"opción": option, **{f"puesto_{p}": round(n, 2) for p, n in sorted(positions.items())}}
        row["promedio"] = round(sum(p * n for p, n in positions.items()) / base, 2)
        row["base_ponderada"] = round(base, 2)
        result.append(row)
    return sorted(result, key=lambda x: x["promedio"])


def frequencies(frame: pd.DataFrame, question: str, kind: str, weights=None) -> list[dict]:
    weights = pd.Series(weights if weights is not None else 1.0, index=frame.index, dtype=float)
    if kind == "matriz de escala":
        counts: dict[str, Counter] = defaultdict(Counter)
        denominators = Counter()
        for idx, raw in frame[question].dropna().items():
            for item, score in parse_matrix_scale(raw):
                counts[item][score] += weights.loc[idx]
                denominators[item] += weights.loc[idx]
        result = []
        for item, scores in counts.items():
            average = sum(score * count for score, count in scores.items()) / max(denominators[item], 1)
            for score, count in sorted(scores.items()):
                result.append({"respuesta": f"{item}: {score}", "ítem": item, "puntaje": score,
                               "cantidad_ponderada": round(count, 2),
                               "porcentaje": round(count / max(denominators[item], 1) * 100, 2),
                               "promedio_ítem": round(average, 2)})
        return result
    if kind == "selección múltiple":
        totals = Counter()
        for idx, raw in frame[question].dropna().items():
            for value in split_answers(raw): totals[value] += weights.loc[idx]
        denominator = max(weights.loc[frame[question].notna()].sum(), 1)
    else:
        totals = Counter()
        for idx, value in frame[question].dropna().astype(str).items(): totals[value] += weights.loc[idx]
        denominator = max(weights.loc[frame[question].notna()].sum(), 1)
    return [{"respuesta": k, "cantidad_ponderada": round(v, 2), "porcentaje": round(v / denominator * 100, 2)} for k, v in totals.most_common()]


def segmentation_slices(project: dict, sample: pd.DataFrame, weights: pd.Series):
    """Total, cortes por cada dimensión y celdas completas de cuota."""
    dims = project.get("quota_settings", {}).get("dimensions", [])
    yield {"nivel_desglose": "Total", "variable_segmentacion": "Total",
           "valor_segmentacion": "Total", "cuota": "Total", **{d: "" for d in dims}}, sample, weights
    for dimension in dims:
        values = list(dict.fromkeys(str(q.get(dimension, "")) for q in project.get("quota_rows", [])))
        for value in values:
            mask = sample[dimension].fillna("").astype(str).eq(value)
            yield {"nivel_desglose": "Variable", "variable_segmentacion": dimension,
                   "valor_segmentacion": value, "cuota": f"{dimension}: {value}",
                   **{d: value if d == dimension else "" for d in dims}}, sample.loc[mask], weights.loc[mask]
    for quota_number, quota in enumerate(project.get("quota_rows", []), 1):
        mask = pd.Series(True, index=sample.index)
        for dimension in dims:
            mask &= sample[dimension].fillna("").astype(str).eq(str(quota.get(dimension, "")))
        yield {"nivel_desglose": "Cuota completa", "variable_segmentacion": " + ".join(dims),
               "valor_segmentacion": " / ".join(str(quota.get(d, "")) for d in dims),
               "cuota": quota_number, **{d: quota.get(d, "") for d in dims}}, sample.loc[mask], weights.loc[mask]


def rankings_by_quota(project: dict, frame: pd.DataFrame) -> list[dict]:
    sample = sample_frame(frame, project)
    chosen = set(project.get('sample_response_ids', []))
    ids = [rid for rid in response_ids(frame) if rid in chosen]
    weights = pd.Series([project.get('sample_weights', {}).get(rid, 1.0) for rid in ids], index=sample.index)
    questions = [x["pregunta"] for x in project.get("dictionary", [])
                 if x.get("incluir") and x.get("tipo") == "ordenamiento/ranking" and x["pregunta"] in sample]
    output = []
    for metadata, quota_frame, quota_weights in segmentation_slices(project, sample, weights):
        for question in questions:
            for result in ranking_table(quota_frame, question, quota_weights):
                output.append({**metadata, "pregunta": question, **result})
    return output


def normalized_text_values(project: dict, sample: pd.DataFrame, question: str, ids: list[str]) -> list[str]:
    """Valor aprobado/normalizado por ID; si falta, conserva la respuesta original."""
    mapping = project.get("text_classifications", {}).get(question, {})
    overrides = project.get("text_response_overrides", {}).get(question, {})
    values = []
    for rid, value in zip(ids, sample[question]):
        original = "" if value is None or bool(pd.isna(value)) else str(value)
        item = overrides.get(rid, mapping.get(original, {}))
        values.append(str(item.get("segment") or original).strip())
    return values


def analysis_by_quota(project: dict, frame: pd.DataFrame, question: str, kind: str) -> dict:
    sample = sample_frame(frame, project)
    chosen = set(project.get('sample_response_ids', []))
    ids = [rid for rid in response_ids(frame) if rid in chosen]
    weights = pd.Series([project.get('sample_weights', {}).get(rid, 1.0) for rid in ids], index=sample.index)
    total = frequencies(sample, question, kind, weights)
    dims = project.get("quota_settings", {}).get("dimensions", [])
    quotas = []
    for quota_number, quota in enumerate(project.get("quota_rows", []), 1):
        mask = pd.Series(True, index=sample.index)
        for dimension in dims:
            mask &= sample[dimension].fillna("").astype(str).eq(str(quota.get(dimension, "")))
        quota_frame, quota_weights = sample.loc[mask], weights.loc[mask]
        valid = quota_frame[question].notna() & quota_frame[question].astype(str).str.strip().ne("")
        quotas.append({"cuota": quota_number, "dimensiones": {d: quota.get(d, "") for d in dims},
                       "personas_reales": int(valid.sum()),
                       "base_ponderada": round(float(quota_weights.loc[valid].sum()), 2),
                       "rows": frequencies(quota_frame, question, kind, quota_weights)})
    return {"question": question, "total": total, "quotas": quotas, "dimensions": dims}


def all_analysis_by_quota_rows(project: dict, frame: pd.DataFrame) -> list[dict]:
    """Tabla larga: dimensiones primero y únicamente combinaciones completas de cuota."""
    output = []
    eligible = {"selección única", "selección múltiple", "escala", "matriz de escala", "demográfica", "texto libre", "ordenamiento/ranking"}
    dims = project.get("quota_settings", {}).get("dimensions", [])
    sample = sample_frame(frame, project)
    chosen = set(project.get('sample_response_ids', []))
    ids = [rid for rid in response_ids(frame) if rid in chosen]
    weights = pd.Series([project.get('sample_weights', {}).get(rid, 1.0) for rid in ids], index=sample.index)
    for item in project.get("dictionary", []):
        question, kind = item.get("pregunta"), item.get("tipo", "")
        if not item.get("incluir") or kind not in eligible or question not in frame:
            continue
        question_sample = sample.copy() if kind == "texto libre" else sample
        analysis_kind = kind
        if kind == "texto libre":
            question_sample[question] = normalized_text_values(project, question_sample, question, ids)
            analysis_kind = "selección única"
        for quota in project.get("quota_rows", []):
            mask = pd.Series(True, index=question_sample.index)
            for dimension in dims:
                mask &= question_sample[dimension].fillna("").astype(str).eq(str(quota.get(dimension, "")))
            segment, segment_weights = question_sample.loc[mask], weights.loc[mask]
            valid = segment[question].notna() & segment[question].astype(str).str.strip().ne("")
            base = round(float(segment_weights.loc[valid].sum()), 2)
            dimensions = {d: quota.get(d, "") for d in dims}
            if kind == "ordenamiento/ranking":
                for row in ranking_table(segment, question, segment_weights):
                    option_base = float(row.get("base_ponderada", 0))
                    for key, amount in row.items():
                        if not key.startswith("puesto_"):
                            continue
                        position = int(key.split("_", 1)[1])
                        output.append({**dimensions, "pregunta": question, "respuesta": row.get("opción"),
                                       "valor": position, "cantidad_ponderada": amount,
                                       "porcentaje": round(float(amount) / option_base * 100, 2) if option_base else 0,
                                       "personas_reales": int(valid.sum()), "base_ponderada": base,
                                       "tipo_pregunta": kind})
            else:
                for row in frequencies(segment, question, analysis_kind, segment_weights):
                    response = row.get("ítem", row.get("respuesta", "")) if kind == "matriz de escala" else row.get("respuesta", "")
                    value = row.get("puntaje", "") if kind == "matriz de escala" else ""
                    output.append({**dimensions, "pregunta": question, "respuesta": response, "valor": value,
                                   "cantidad_ponderada": row.get("cantidad_ponderada", 0),
                                   "porcentaje": row.get("porcentaje", 0),
                                   "personas_reales": int(valid.sum()), "base_ponderada": base,
                                   "tipo_pregunta": kind})
    return output


def export_workbook(project: dict, frame: pd.DataFrame) -> bytes:
    all_ids = response_ids(frame)
    original = frame.copy()
    original['ID de respuesta'] = all_ids
    clean = sample_frame(frame, project)
    chosen = set(project.get('sample_response_ids', []))
    review_ids = [rid for rid in all_ids if rid in chosen]
    clean['peso_muestral'] = [project.get('sample_weights', {}).get(rid, 1.0) for rid in review_ids]
    clean['ID de respuesta'] = review_ids
    clean['ID de revisión'] = review_ids
    text_questions = set(project.get("text_classifications", {})) | set(project.get("text_response_overrides", {}))
    text_questions.update(item.get("pregunta") for item in project.get("dictionary", [])
                          if item.get("incluir") and item.get("tipo") == "texto libre")
    for question in text_questions:
        if question in clean:
            clean[question] = normalized_text_values(project, clean, question, review_ids)
    for item in project.get("dictionary", []):
        question = item.get("pregunta")
        if not item.get("incluir") or question not in clean:
            continue
        if item.get("tipo") == "texto libre":
            # Ya quedó reemplazada por su único valor canónico.
            continue
        elif item.get("tipo") == "ordenamiento/ranking":
            parsed = clean[question].map(lambda x: dict(parse_ranking(x)))
            for option in sorted({o for row in parsed for o in row}):
                clean[f"{question} [puesto: {option}]"] = parsed.map(lambda x: x.get(option))
            clean.drop(columns=[question], inplace=True)
        elif item.get("tipo") == "matriz de escala":
            parsed = clean[question].map(lambda x: dict(parse_matrix_scale(x)))
            for scale_item in sorted({name for row in parsed for name in row}):
                clean[f"{question} [escala: {scale_item}]"] = parsed.map(lambda x: x.get(scale_item))
            clean.drop(columns=[question], inplace=True)
    # Copia analítica con las estructuras originales necesarias para calcular matrices
    # y rankings; el texto libre sí usa su único valor canónico.
    analysis_clean = sample_frame(frame, project)
    analysis_clean['peso_muestral'] = [project.get('sample_weights', {}).get(rid, 1.0) for rid in review_ids]
    for item in project.get("dictionary", []):
        question = item.get("pregunta")
        if item.get("incluir") and item.get("tipo") == "texto libre" and question in analysis_clean:
            analysis_clean[question] = normalized_text_values(project, analysis_clean, question, review_ids)
    # La base limpia conserva una fila por persona. La ponderada materializa el ajuste
    # con réplicas aleatorias, sin hacerlas pasar por nuevos encuestados.
    expanded_ids = project.get('sample_expanded_response_ids')
    if not expanded_ids and project.get('sample_weighted'):
        expanded_ids = expanded_sample_ids(frame, project, review_ids, int(project.get('sample_seed', 2026)))
    expanded_ids = expanded_ids or review_ids
    by_id = {rid: row for rid, row in zip(review_ids, clean.to_dict("records"))}
    occurrences = Counter()
    expanded_rows = []
    for rid in expanded_ids:
        if rid not in by_id:
            continue
        occurrences[rid] += 1
        row = dict(by_id[rid])
        row['ID de respuesta original'] = rid
        row['ID de fila representativa'] = f'{rid}-{occurrences[rid]}'
        row['es_réplica_estadística'] = occurrences[rid] > 1
        row['peso_muestral'] = 1.0
        expanded_rows.append(row)

    wb = Workbook()
    ows = wb.active; ows.title = "Base original"
    append_sheet(ows, original.to_dict("records"))
    ws = wb.create_sheet("Base limpia")
    append_sheet(ws, clean.to_dict("records"))
    ews = wb.create_sheet("Base ponderada")
    append_sheet(ews, expanded_rows)
    qws = wb.create_sheet("Cuotas"); append_sheet(qws, project.get("quota_rows", []))
    aws = wb.create_sheet("Respuestas por cuota"); append_sheet(aws, all_analysis_by_quota_rows(project, frame))
    rws = wb.create_sheet("Rankings por cuota"); append_sheet(rws, rankings_by_quota(project, frame))
    gws = wb.create_sheet("Gráficos")
    dictionary = {x["pregunta"]: x for x in project.get("dictionary", [])}
    row = 1
    for question, item in dictionary.items():
        if question not in analysis_clean or not item.get("incluir") or item.get("tipo") in {"identificador", "técnica", "ordenamiento/ranking"}:
            continue
        kind = "selección única" if item.get("tipo") == "texto libre" else item.get("tipo", "")
        table = frequencies(analysis_clean, question, kind, analysis_clean['peso_muestral'])[:30]
        if not table: continue
        gws.cell(row, 1, question).font = Font(bold=True, size=12)
        start = row + 1
        for i, x in enumerate(table, start):
            gws.cell(i, 1, x["respuesta"]); gws.cell(i, 2, x["cantidad_ponderada"])
        chart = BarChart(); chart.title = question[:80]
        chart.add_data(Reference(gws, min_col=2, min_row=start, max_row=start + len(table) - 1))
        chart.set_categories(Reference(gws, min_col=1, min_row=start, max_row=start + len(table) - 1))
        chart.height = 7; chart.width = 14
        gws.add_chart(chart, f"D{start}")
        row += max(len(table) + 3, 16)
    sgws = wb.create_sheet("Gráficos segmentados")
    segment_weights = pd.Series(clean['peso_muestral'].tolist(), index=clean.index, dtype=float)
    row = 1
    for question, item in dictionary.items():
        if question not in analysis_clean or not item.get("incluir") or item.get("tipo") in {"identificador", "técnica", "ordenamiento/ranking"}:
            continue
        kind = "selección única" if item.get("tipo") == "texto libre" else item.get("tipo", "")
        for metadata, segment, weights in segmentation_slices(project, analysis_clean, segment_weights):
            if metadata["nivel_desglose"] != "Variable":
                continue
            results = frequencies(segment, question, kind, weights)[:30]
            if not results:
                continue
            title = f'{question} · {metadata["variable_segmentacion"]}: {metadata["valor_segmentacion"]}'
            sgws.cell(row, 1, title).font = Font(bold=True, size=12)
            sgws.cell(row + 1, 1, "Respuesta"); sgws.cell(row + 1, 2, "Cantidad ponderada"); sgws.cell(row + 1, 3, "Porcentaje")
            for cell in sgws[row + 1]:
                cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17324D")
            start = row + 2
            for i, result in enumerate(results, start):
                sgws.cell(i, 1, result["respuesta"]); sgws.cell(i, 2, result["cantidad_ponderada"]); sgws.cell(i, 3, result["porcentaje"])
            chart = BarChart(); chart.title = title[:120]
            chart.add_data(Reference(sgws, min_col=2, min_row=start, max_row=start + len(results) - 1))
            chart.set_categories(Reference(sgws, min_col=1, min_row=start, max_row=start + len(results) - 1))
            chart.height = 7; chart.width = 14
            sgws.add_chart(chart, f"E{start}")
            row += max(len(results) + 4, 16)
    sws = wb.create_sheet("Cruces sugeridos")
    suggestions = suggest_crosses(clean, project.get("dictionary", []))
    append_sheet(sws, suggestions)
    output = io.BytesIO(); wb.save(output); return output.getvalue()


def storage_rows(project: dict, frame: pd.DataFrame):
    """Construye las tres capas persistentes sin guardar secretos ni depender del XLSX."""
    all_ids = response_ids(frame)
    raw = [{"project_id": project["id"], "response_id": rid, "payload": clean_json(row)}
           for rid, row in zip(all_ids, frame.to_dict("records"))]
    selected = set(project.get("sample_response_ids", []))
    sample = sample_frame(frame, project)
    selected_ids = [rid for rid in all_ids if rid in selected]
    processed_frame = sample.copy()
    for item in project.get("dictionary", []):
        question = item.get("pregunta")
        if item.get("incluir") and item.get("tipo") == "texto libre" and question in processed_frame:
            processed_frame[question] = normalized_text_values(project, processed_frame, question, selected_ids)
    by_id = {rid: clean_json(row) for rid, row in zip(selected_ids, processed_frame.to_dict("records"))}
    expanded = project.get("sample_expanded_response_ids") or selected_ids
    occurrences = Counter()
    processed = []
    for rid in expanded:
        if rid not in by_id:
            continue
        occurrences[rid] += 1
        processed.append({"project_id": project["id"], "response_id": rid, "payload": by_id[rid],
                          "sample_weight": project.get("sample_weights", {}).get(rid, 1.0),
                          "is_statistical_replica": occurrences[rid] > 1,
                          "replica_number": occurrences[rid]})
    dims = project.get("quota_settings", {}).get("dimensions", [])
    results = []
    for row in all_analysis_by_quota_rows(project, frame):
        results.append({"project_id": project["id"], "dimensions": {d: row.get(d, "") for d in dims},
                        "question": row.get("pregunta", ""), "answer": row.get("respuesta"),
                        "value": str(row.get("valor", "")) if row.get("valor", "") != "" else None,
                        "weighted_count": row.get("cantidad_ponderada"), "percentage": row.get("porcentaje"),
                        "real_respondents": row.get("personas_reales"), "weighted_base": row.get("base_ponderada"),
                        "question_type": row.get("tipo_pregunta")})
    return raw, processed, results


def append_sheet(ws, rows: list[dict]) -> None:
    if not rows: return
    headers = list(dict.fromkeys(k for row in rows for k in row))
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17324D")
    for row in rows:
        ws.append([row.get(h) if not isinstance(row.get(h), (dict, list)) else json.dumps(row.get(h), ensure_ascii=False) for h in headers])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions


def suggest_crosses(frame: pd.DataFrame, dictionary: list[dict]) -> list[dict]:
    demographics = [x["pregunta"] for x in dictionary if x.get("incluir") and x.get("tipo") == "demográfica" and x["pregunta"] in frame]
    outcomes = [x["pregunta"] for x in dictionary if x.get("incluir") and x.get("tipo") in {"selección única", "escala", "demográfica"} and 2 <= x.get("valores_unicos", 0) <= 20 and x["pregunta"] in frame]
    result = []
    for a in demographics:
        for b in outcomes:
            if a == b: continue
            table = pd.crosstab(frame[a], frame[b])
            if table.shape[0] < 2 or table.shape[1] < 2 or table.values.sum() < 30: continue
            observed = table.to_numpy(float); n = observed.sum(); expected = observed.sum(1)[:, None] * observed.sum(0)[None, :] / n
            chi = (((observed - expected) ** 2 / expected)[expected > 0]).sum()
            score = math.sqrt((chi / n) / max(min(table.shape) - 1, 1))
            result.append({"segmentar_por": a, "respuesta_a_comparar": b, "fuerza": round(score, 3), "base": int(n)})
    return sorted(result, key=lambda x: x["fuerza"], reverse=True)[:20]


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def send_json(self, value: Any, status=200):
        raw = json.dumps(clean_json(value), ensure_ascii=False).encode()
        self.send_response(status); self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw))); self.end_headers(); self.wfile.write(raw)

    def body(self) -> dict:
        size = int(self.headers.get("Content-Length", 0)); return json.loads(self.rfile.read(size) or b"{}")

    def do_GET(self):
        try:
            path, _, query = self.path.partition("?")
            if path.startswith("/api/job/"):
                job_id = path.rsplit("/", 1)[-1]
                job = read_job(job_id)
                if not job:
                    return self.send_json({"error": "Tarea inexistente"}, 404)
                if job.get("status") == "running" and time.time() - float(job.get("updated_at", 0)) > 180:
                    job.update({"status": "interrupted", "phase": "La instancia se reinició; los lotes ya guardados se conservan",
                                "error": "La ejecución se interrumpió por reinicio del servidor. Volvé a procesar para continuar desde los pendientes."})
                return self.send_json(job)
            if path == "/api/text-jobs":
                jobs = {job_id: dict(job) for job_id, job in TEXT_JOBS.items()}
                for summary in list_projects():
                    try:
                        project = read_project(summary["id"]); job_id = project.get("last_text_job_id")
                        if job_id and job_id not in jobs:
                            job = read_job(job_id)
                            if job: jobs[job_id] = job
                    except Exception:
                        continue
                return self.send_json(list(jobs.values()))
            if path == "/api/projects":
                hydrate_cloud_projects()
                return self.send_json(list_projects())
            if path == "/api/cloud-projects": return self.send_json(gcs_projects.list_cloud_projects())
            if path.startswith("/api/project/"):
                parts = path.split("/"); pid = parts[3]; project = read_project(pid)
                if len(parts) == 4: return self.send_json(project)
                if parts[4] == "cloud-status":
                    return self.send_json(CLOUD_SYNC_STATUS.get(pid, {
                        "ok": None, "pending": pid in CLOUD_SYNC_TIMERS,
                        "message": "El primer checkpoint todavía está pendiente."
                    }))
                if parts[4] == "text-review":
                    question = urllib.parse.parse_qs(query).get("question", [""])[0]
                    mapping = project.get("text_classifications", {}).get(question, {})
                    return self.send_json([{"respuesta": k, **v} for k, v in mapping.items()])
                frame = load_source(project)
                if parts[4] == 'external-review':
                    question = resolve_column(frame, urllib.parse.parse_qs(query).get('question', [''])[0])
                    sample = sample_frame(frame, project)
                    ids = [rid for rid in response_ids(frame) if rid in set(project.get('sample_response_ids', []))]
                    rows = external_review.sheet_rows(project, sample, question, ids)
                    cfg = project.get('text_question_settings', {}).get(question, {})
                    return self.send_json({'rows': rows, 'tsv': external_review.tsv(rows),
                        'prompt': external_review.review_prompt(question, cfg.get('mode', 'semantic'), cfg.get('instructions', ''))})
                if parts[4] == 'cross-options':
                    sample=sample_frame(frame,project); ids=[rid for rid in response_ids(frame) if rid in set(project.get('sample_response_ids',[]))]
                    data,kinds=cross_analysis.dataset(project,sample,ids,split_answers)
                    return self.send_json({'values':cross_analysis.values(data,kinds,split_answers),'kinds':kinds})
                if parts[4] == "workflow-status":
                    text_status = text_workflow_status(frame, project)
                    ranking_questions = [x["pregunta"] for x in project.get("dictionary", [])
                                         if x.get("incluir") and x.get("tipo") == "ordenamiento/ranking"]
                    processed_rankings = set(project.get("selected_ranking_questions", []))
                    return self.send_json({"text": text_status,
                        "rankings_complete": all(q in processed_rankings for q in ranking_questions),
                        "ranking_total": len(ranking_questions),
                        "ranking_processed": sum(q in processed_rankings for q in ranking_questions)})
                if parts[4] == "source": return self.send_json({"rows": len(frame), "columns": list(frame.columns),
                    "dictionary": merge_dictionary(project.get("dictionary", []), build_dictionary(frame)),
                    "refreshed_at": project.get("source_refreshed_at")})
                if parts[4] == "quota-options": return self.send_json({c: sorted(frame[c].dropna().astype(str).unique().tolist()) for c in frame.columns if 1 < frame[c].nunique() <= 80})
                if parts[4] == "sample": return self.send_json({"rows": len(sample_frame(frame, project)), "target": project.get("quota_settings", {}).get("target_total", 0)})
                if parts[4] == "text-review":
                    question = urllib.parse.unquote(query.split("question=", 1)[1]) if "question=" in query else ""
                    mapping = project.get("text_classifications", {}).get(question, {})
                    return self.send_json([{"respuesta": k, **v} for k, v in mapping.items()])
                if parts[4] == "ranking":
                    params = urllib.parse.parse_qs(query); question = resolve_column(frame, params.get("question", [""])[0])
                    base = sample_frame(frame, project)
                    chosen = set(project.get('sample_response_ids', []))
                    ids = [rid for rid in response_ids(frame) if rid in chosen]
                    weights = pd.Series([project.get('sample_weights', {}).get(rid, 1.0) for rid in ids], index=base.index)
                    quota_index = params.get("quota", [""])[0]
                    label = "Todas las cuotas"
                    if quota_index != "":
                        index = int(quota_index)
                        rows = project.get("quota_rows", [])
                        if index < 0 or index >= len(rows):
                            raise ValueError("La cuota seleccionada ya no existe.")
                        quota = rows[index]; dims = project.get("quota_settings", {}).get("dimensions", [])
                        mask = pd.Series(True, index=base.index)
                        for dimension in dims:
                            mask &= base[dimension].fillna("").astype(str).eq(str(quota.get(dimension, "")))
                        base = base.loc[mask]; weights = weights.loc[mask]
                        label = " / ".join(str(quota.get(d, "")) for d in dims)
                    valid = base[question].notna() & base[question].astype(str).str.strip().ne("")
                    return self.send_json({"rows": ranking_table(base, question, weights), "quota": label,
                        "personas_reales": int(valid.sum()), "base_ponderada": round(float(weights.loc[valid].sum()), 2)})
                if parts[4] == "analysis":
                    params = urllib.parse.parse_qs(query); question = resolve_column(frame, params.get("question", [""])[0])
                    dictionary = {x["pregunta"]: x for x in project.get("dictionary", [])}
                    base = sample_frame(frame, project)
                    ids = [rid for rid in response_ids(frame) if rid in set(project.get('sample_response_ids', []))]
                    weights = pd.Series([project.get('sample_weights', {}).get(rid, 1.0) for rid in ids], index=base.index)
                    return self.send_json(frequencies(base, question, dictionary.get(question, {}).get("tipo", ""), weights))
                if parts[4] == "analysis-by-quota":
                    params = urllib.parse.parse_qs(query); question = resolve_column(frame, params.get("question", [""])[0])
                    dictionary = {x["pregunta"]: x for x in project.get("dictionary", [])}
                    return self.send_json(analysis_by_quota(project, frame, question,
                        dictionary.get(question, {}).get("tipo", "")))
                if parts[4] == "suggestions": return self.send_json(suggest_crosses(sample_frame(frame, project), project.get("dictionary", [])))
                if parts[4] == "export":
                    raw = export_workbook(project, frame); self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", f'attachment; filename="{project.get("name","encuesta")}.xlsx"')
                    self.send_header("Content-Length", str(len(raw))); self.end_headers(); return self.wfile.write(raw)
            if path == "/" or path == "/index.html": return self.send_file(STATIC_DIR / "index.html", "text/html; charset=utf-8")
            if path.startswith("/static/"):
                file = STATIC_DIR / path.removeprefix("/static/")
                mime = "text/css" if file.suffix == ".css" else "application/javascript"
                return self.send_file(file, mime)
            self.send_error(404)
        except Exception as exc:
            self.send_json({"error": str(exc), "details": traceback.format_exc()}, 400)

    def send_file(self, path: Path, mime: str):
        raw = path.read_bytes(); self.send_response(200); self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(raw))); self.end_headers(); self.wfile.write(raw)

    def do_POST(self):
        try:
            path = self.path.split("?", 1)[0]; data = self.body()
            if path == "/api/projects":
                project = {"id": uuid.uuid4().hex, "version": 2, "name": data.get("name", "Nuevo proyecto").strip(), "description": data.get("description", ""), "source_url": "", "quota_settings": {}, "quota_rows": [], "sample_response_ids": [], "dictionary": [], "text_classifications": {}, "text_codebooks": {}, "created_at": datetime.now().isoformat(timespec="seconds")}
                save_project(project); return self.send_json(project, 201)
            if path == "/api/cloud-projects/restore":
                project = gcs_projects.restore(data.get("prefix", ""), PROJECTS_DIR, SNAPSHOTS_DIR)
                DATA_CACHE.pop(project["id"], None)
                return self.send_json(project)
            parts = path.split("/"); pid = parts[3]; action = parts[4]; project = read_project(pid)
            if action == "save":
                sheet_changed = (
                    "source_sheet_gid" in data
                    and str(data.get("source_sheet_gid", "")) != str(project.get("source_sheet_gid", ""))
                )
                changed_questions = reconcile_changed_questions(project, data["dictionary"]) if "dictionary" in data else []
                for key in ["name", "description", "source_url", "source_sheet_gid", "source_sheet_name", "dictionary", "quota_settings", "quota_rows", "ai", "selected_ranking_questions"]:
                    if key in data: project[key] = data[key]
                if sheet_changed:
                    project["sample_response_ids"] = []
                    project["text_classifications"] = {}
                    project["text_codebooks"] = {}
                    project["text_processing_modes"] = {}
                    project["selected_ranking_questions"] = []
                if changed_questions:
                    project["last_changed_questions"] = changed_questions
                save_project(project); return self.send_json(project)
            if action == "discover-sheets":
                return self.send_json(list_google_sheet_tabs(data.get("url") or project.get("source_url", "")))
            if action == "test-ai":
                return self.send_json(test_gemini(data))
            if action == "text-settings":
                questions, settings = text_selection(project, data)
                project["selected_text_questions"] = questions
                project["text_question_settings"] = settings
                save_project(project)
                return self.send_json({"saved": True})
            if action == "stop-text":
                with LOCK:
                    for job in TEXT_JOBS.values():
                        if job["project_id"] == pid and job["status"] == "running":
                            job["stop_requested"] = True
                return self.send_json({"message": "Se detendrá después de la llamada actual."})
            if action == "save-cloud":
                workbook = None
                if project.get("sample_response_ids"):
                    workbook = export_workbook(project, load_source(project))
                result = gcs_projects.save_bundle(
                    project, project_path(pid), SNAPSHOTS_DIR / f"{pid}.jsonl", workbook
                )
                CLOUD_SYNC_STATUS[pid] = {
                    "ok": True, "saved_at": datetime.now().isoformat(timespec="seconds"), **result
                }
                return self.send_json(result)
            frame = load_source(project, refresh=action == "refresh-source")
            if action == "sync-storage":
                raw_rows, processed_rows, result_rows = storage_rows(project, frame)
                return self.send_json(storage_sync.sync(project, raw_rows, processed_rows, result_rows))
            if action == 'cross-analysis':
                sample=sample_frame(frame,project); ids=[rid for rid in response_ids(frame) if rid in set(project.get('sample_response_ids',[]))]
                analysis,kinds=cross_analysis.dataset(project,sample,ids,split_answers)
                return self.send_json(cross_analysis.analyze(analysis,kinds,data['primary'],data.get('breakdowns',[]),data.get('filters',[]),split_answers))
            if action in ('external-preview', 'external-apply'):
                with LOCK:
                    latest = read_project(pid)
                    question = resolve_column(frame, data['question'])
                    sample = sample_frame(frame, latest)
                    ids = [rid for rid in response_ids(frame) if rid in set(latest.get('sample_response_ids', []))]
                    rows = external_review.sheet_rows(latest, sample, question, ids)
                    comparison = external_review.compare(latest, question, rows, data['tsv'])
                    if action == 'external-apply':
                        if data.get('revision') != comparison['revision']:
                            raise ValueError('La muestra o los resultados cambiaron. Volvé a previsualizar antes de guardar.')
                        external_review.apply_changes(latest, question, comparison['changes'])
                        save_project(latest)
                    return self.send_json(comparison)
            if action == "text-prompt":
                return self.send_json(text_prompt_preview(project, frame, data))
            if action == "refresh-source": return self.send_json({"rows": len(frame), "columns": list(frame.columns), "dictionary": project["dictionary"], "refreshed_at": project.get("source_refreshed_at")})
            if action == "quota-template": return self.send_json(make_quota_rows(frame, data["dimensions"], int(data["target_total"]), data["input_mode"]))
            if action == "save-quotas":
                status = quota_status(frame, data["dimensions"], data["rows"], data["input_mode"], int(data["target_total"]))
                project["quota_settings"] = {"dimensions": data["dimensions"], "sort_dimensions": data.get("sort_dimensions", data["dimensions"]), "input_mode": data["input_mode"], "target_total": int(data["target_total"])}
                project["quota_rows"] = data["rows"]; project["sample_response_ids"] = []
                project["sample_expanded_response_ids"] = []; project["sample_weights"] = {}; save_project(project)
                return self.send_json(status)
            if action == "quota-status": return self.send_json(quota_status(frame, data["dimensions"], data["rows"], data["input_mode"], int(data["target_total"])))
            if action == "balance":
                weighted = bool(data.get('weighted'))
                ids, count, weights = balanced_sample(frame, project, int(data.get("seed", 2026)), weighted)
                project["sample_response_ids"] = ids; project['sample_weights'] = weights; project['sample_weighted'] = weighted
                project["sample_expanded_response_ids"] = expanded_sample_ids(frame, project, ids, int(data.get("seed", 2026)))
                project["sample_seed"] = int(data.get("seed", 2026)); project["sample_saved_at"] = datetime.now().isoformat(timespec="seconds"); save_project(project)
                return self.send_json({"selected": count, "discarded": len(frame) - count, 'weighted': weighted,
                    'weighted_total': round(sum(weights.values()), 2), 'expanded_total': len(project["sample_expanded_response_ids"]),
                    'max_weight': max(weights.values(), default=0)})
            if action == "process-text":
                return self.send_json(start_text_job(pid, data), 202)
            if action == "review-text-by-id":
                with LOCK:
                    latest = read_project(pid)
                    question = resolve_column(frame, data["question"])
                    sample = sample_frame(frame, latest)
                    sample_ids = [rid for rid in response_ids(frame) if rid in set(latest.get("sample_response_ids", []))]
                    allowed = {row["id"]: row for row in external_review.sheet_rows(latest, sample, question, sample_ids)}
                    overrides = latest.setdefault("text_response_overrides", {}).setdefault(question, {})
                    mode = latest.get("text_question_settings", {}).get(question, {}).get("mode") or latest.get("text_processing_modes", {}).get(question)
                    known, _ = brand_memory(latest)
                    saved = 0
                    for item in data.get("rows", []):
                        rid = str(item.get("id", ""))
                        if rid not in allowed:
                            raise ValueError(f"ID desconocido para esta pregunta: {rid}")
                        segment = str(item.get("segment", "")).strip()
                        if allowed[rid]["original"].strip() and not segment:
                            raise ValueError(f"La respuesta procesada del ID {rid} no puede quedar vacía.")
                        if mode == "brands" and segment:
                            segment = normalize_brand(segment, known)
                            known.extend(segment.split(";"))
                        overrides[rid] = {"segment": segment, "original": allowed[rid]["original"],
                                          "approved": bool(item.get("approved", False)), "reviewed": True,
                                          "source": "inline_manual"}
                        saved += 1
                    save_project(latest)
                return self.send_json({"saved": saved})
            if action == "finalize-text-review":
                with LOCK:
                    latest = read_project(pid)
                    sample = sample_frame(frame, latest)
                    sample_ids = [rid for rid in response_ids(frame) if rid in set(latest.get("sample_response_ids", []))]
                    questions = latest.get("selected_text_questions") or [
                        row["pregunta"] for row in latest.get("dictionary", [])
                        if row.get("incluir") and row.get("tipo") == "texto libre"
                    ]
                    approved = 0
                    for requested in questions:
                        question = resolve_column(sample, requested)
                        mapping = latest.get("text_classifications", {}).get(question, {})
                        overrides = latest.setdefault("text_response_overrides", {}).setdefault(question, {})
                        for rid, raw in zip(sample_ids, sample[question]):
                            original = "" if raw is None or bool(pd.isna(raw)) else str(raw)
                            if not original.strip():
                                continue
                            current = overrides.get(rid, mapping.get(original, {}))
                            segment = str(current.get("segment", "")).strip()
                            if not segment:
                                raise ValueError(f"Todavía falta procesar el ID {rid} de la pregunta: {question}")
                            overrides[rid] = {**current, "segment": segment, "original": original,
                                              "approved": True, "reviewed": True,
                                              "source": current.get("source") or "gemini_confirmed"}
                            approved += 1
                    save_project(latest)
                return self.send_json({"approved": approved, "questions": len(questions)})
            if action == "review-text":
                with LOCK:
                    latest = read_project(pid)
                    mapping = latest.setdefault("text_classifications", {}).setdefault(data["question"], {})
                    mode = latest.get("text_processing_modes", {}).get(data["question"]) or latest.get("text_question_settings", {}).get(data["question"], {}).get("mode")
                    known, _ = brand_memory(latest)
                    for item in data["rows"]:
                        if item["respuesta"] in mapping:
                            segment = normalize_brand(item["segment"], known) if mode == "brands" else item["segment"]
                            if mode == "brands": known.extend(segment.split(';'))
                            mapping[item["respuesta"]].update({"segment": segment, "approved": bool(item.get("approved")), "reviewed": True, "needs_review": brand_key(segment) == "requiererevisión"})
                    save_project(latest)
                return self.send_json({"saved": len(data["rows"])})
            self.send_error(404)
        except Exception as exc:
            self.send_json({"error": str(exc), "details": traceback.format_exc()}, 400)

    def do_DELETE(self):
        try:
            path = self.path.split("?", 1)[0]
            match = re.fullmatch(r"/api/project/([A-Za-z0-9_-]+)", path)
            if not match:
                return self.send_error(404)
            return self.send_json(delete_project(match.group(1)))
        except FileNotFoundError as exc:
            self.send_json({"error": str(exc)}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc), "details": traceback.format_exc()}, 400)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--port", type=int, default=int(os.getenv("PORT", "8765"))); parser.add_argument("--host", default=os.getenv("HOST", "127.0.0.1")); parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args(); server = ThreadingHTTPServer((args.host, args.port), Handler)
    url = f"http://{args.host}:{args.port}"
    print(f"D1ana Wizard Tool disponible en {url}")
    if not args.no_browser and args.host in {"127.0.0.1", "localhost"}: threading.Timer(1, lambda: webbrowser.open(url)).start()
    try: server.serve_forever()
    except KeyboardInterrupt: pass


if __name__ == "__main__": main()
