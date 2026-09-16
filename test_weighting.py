import unittest
import io
import pandas as pd
from openpyxl import load_workbook
import server

class WeightingTests(unittest.TestCase):
    def project(self, rows):
        return {'quota_settings': {'dimensions':['zona'],'input_mode':'Cantidades','target_total':10}, 'quota_rows':rows}
    def test_shortfall_weights_to_targets(self):
        frame=pd.DataFrame({'id':range(8),'zona':['A']*3+['B']*5})
        ids,count,weights=server.balanced_sample(frame,self.project([{'zona':'A','cantidad':5},{'zona':'B','cantidad':5}]),7,True)
        self.assertEqual(count,8);self.assertAlmostEqual(sum(weights.values()),10)
        self.assertEqual(sorted(set(weights.values())),[1.0,round(5/3,8)])
    def test_zero_cell_rejected(self):
        frame=pd.DataFrame({'zona':['B']*5})
        with self.assertRaisesRegex(ValueError,'sin respuestas'):
            server.balanced_sample(frame,self.project([{'zona':'A','cantidad':5},{'zona':'B','cantidad':5}]),7,True)
    def test_weighted_frequency(self):
        frame=pd.DataFrame({'Q':['Sí','No']})
        result=server.frequencies(frame,'Q','selección única',pd.Series([3.,1.]))
        self.assertEqual(result[0]['porcentaje'],75)

    def test_expanded_sample_reaches_each_quota_target(self):
        frame=pd.DataFrame({'id':range(8),'zona':['A']*3+['B']*5})
        project=self.project([{'zona':'A','cantidad':5},{'zona':'B','cantidad':5}])
        ids,_,_=server.balanced_sample(frame,project,7,True)
        expanded=server.expanded_sample_ids(frame,project,ids,7)
        by_id=dict(zip(server.response_ids(frame),frame['zona']))
        self.assertEqual(len(expanded),10)
        self.assertEqual([by_id[x] for x in expanded].count('A'),5)
        self.assertEqual([by_id[x] for x in expanded].count('B'),5)

    def test_export_contains_unique_weights_and_expanded_rows(self):
        frame=pd.DataFrame({'id':range(8),'zona':['A']*3+['B']*5,'Q':['Sí']*8})
        project={**self.project([{'zona':'A','cantidad':5},{'zona':'B','cantidad':5}]),
                 'dictionary':[{'pregunta':'Q','tipo':'selección única','incluir':True}],
                 'selected_ranking_questions':[]}
        ids,_,weights=server.balanced_sample(frame,project,7,True)
        project.update(sample_response_ids=ids,sample_weights=weights,sample_weighted=True,sample_seed=7,
                       sample_expanded_response_ids=server.expanded_sample_ids(frame,project,ids,7))
        workbook=load_workbook(io.BytesIO(server.export_workbook(project,frame)),data_only=True)
        self.assertEqual(workbook['Base original'].max_row-1,8)
        self.assertEqual(workbook['Base limpia'].max_row-1,8)
        self.assertEqual(workbook['Base ponderada'].max_row-1,10)
        self.assertIn('Respuestas por cuota',workbook.sheetnames)
        self.assertIn('Rankings por cuota',workbook.sheetnames)
        self.assertIn('Gráficos segmentados',workbook.sheetnames)
        headers=[c.value for c in workbook['Base limpia'][1]]
        self.assertIn('peso_muestral',headers)

    def test_dictionary_change_preserves_previous_processing(self):
        project={'dictionary':[{'pregunta':'A','tipo':'texto libre','incluir':True},
                               {'pregunta':'B','tipo':'texto libre','incluir':True}],
                 'text_classifications':{'A':{'x':{}},'B':{'y':{}}}}
        changed=server.reconcile_changed_questions(project,[
            {'pregunta':'A','tipo':'ordenamiento/ranking','incluir':True},
            {'pregunta':'B','tipo':'texto libre','incluir':True}])
        self.assertEqual(changed,['A'])
        self.assertIn('A',project['text_classifications'])
        self.assertIn('B',project['text_classifications'])

    def test_rankings_are_aggregated_by_quota_with_weights(self):
        frame=pd.DataFrame({'response id':['1','2','3'],'zona':['A','A','B'],
                            'Ranking':['1. Toyota; 2. Ford','1. Ford; 2. Toyota','1. Toyota; 2. Ford']})
        project={'quota_settings':{'dimensions':['zona'],'input_mode':'Cantidades','target_total':5},
                 'quota_rows':[{'zona':'A','cantidad':4},{'zona':'B','cantidad':1}],
                 'sample_response_ids':['1','2','3'],'sample_weights':{'1':2.0,'2':2.0,'3':1.0},
                 'dictionary':[{'pregunta':'Ranking','tipo':'ordenamiento/ranking','incluir':True}]}
        rows=server.rankings_by_quota(project,frame)
        a_toyota=next(r for r in rows if r['zona']=='A' and r['opción']=='Toyota')
        self.assertEqual(a_toyota['puesto_1'],2.0)
        self.assertEqual(a_toyota['puesto_2'],2.0)
        self.assertEqual(a_toyota['base_ponderada'],4.0)

    def test_analysis_returns_total_and_every_quota(self):
        frame=pd.DataFrame({'response id':['1','2','3'],'zona':['A','A','B'],'Elección':['Sí','No','Sí']})
        project={'quota_settings':{'dimensions':['zona'],'input_mode':'Cantidades','target_total':5},
                 'quota_rows':[{'zona':'A','cantidad':4},{'zona':'B','cantidad':1}],
                 'sample_response_ids':['1','2','3'],'sample_weights':{'1':2.0,'2':2.0,'3':1.0}}
        result=server.analysis_by_quota(project,frame,'Elección','selección única')
        self.assertEqual(len(result['quotas']),2)
        self.assertEqual(sum(x['cantidad_ponderada'] for x in result['total']),5)
        self.assertEqual(result['quotas'][0]['base_ponderada'],4)

    def test_export_analysis_includes_each_dimension_and_complete_quota(self):
        frame=pd.DataFrame({'response id':['1','2','3','4'],
                            'género':['Mujer','Hombre','Mujer','Hombre'],
                            'edad':['18-29','18-29','30-44','30-44'],
                            'región':['Norte','Norte','Sur','Sur'],
                            'Elección':['Sí','No','Sí','No']})
        quota_rows=[]
        for genero,edad,region in zip(frame['género'],frame['edad'],frame['región']):
            quota_rows.append({'género':genero,'edad':edad,'región':region,'cantidad':1})
        project={'quota_settings':{'dimensions':['género','edad','región'],'input_mode':'Cantidades','target_total':4},
                 'quota_rows':quota_rows,'sample_response_ids':['1','2','3','4'],
                 'sample_weights':{str(i):1.0 for i in range(1,5)},
                 'dictionary':[{'pregunta':'Elección','tipo':'selección única','incluir':True}]}
        rows=server.all_analysis_by_quota_rows(project,frame)
        self.assertTrue(rows)
        self.assertTrue(all(set(('género','edad','región','pregunta','respuesta','valor')).issubset(r) for r in rows))
        self.assertTrue(all('nivel_desglose' not in r for r in rows))

    def test_results_by_quota_include_normalized_text_and_rankings(self):
        frame=pd.DataFrame({'response id':['1','2'],'zona':['A','B'],
                            'Opinión':['muy caro','muy caro'],'Ranking':['1. Toyota; 2. Ford','1. Ford; 2. Toyota']})
        project={'quota_settings':{'dimensions':['zona'],'input_mode':'Cantidades','target_total':2},
                 'quota_rows':[{'zona':'A','cantidad':1},{'zona':'B','cantidad':1}],
                 'sample_response_ids':['1','2'],'sample_weights':{'1':1.0,'2':1.0},
                 'dictionary':[{'pregunta':'Opinión','tipo':'texto libre','incluir':True},
                               {'pregunta':'Ranking','tipo':'ordenamiento/ranking','incluir':True}],
                 'text_classifications':{'Opinión':{'muy caro':{'segment':'Precio'}}}}
        rows=server.all_analysis_by_quota_rows(project,frame)
        self.assertTrue(any(r['pregunta']=='Opinión' and r['respuesta']=='Precio' for r in rows))
        self.assertTrue(any(r['pregunta']=='Ranking' and r['respuesta']=='Toyota' and
                            r['valor']==1 for r in rows))

    def test_quota_template_builds_full_cartesian_product(self):
        frame=pd.DataFrame({'zona':['Norte','Norte','Sur'],
                            'género':['Mujer','Varón','Varón'],
                            'edad':['18-29','30-44','18-29']})
        rows=server.make_quota_rows(frame,['zona','género','edad'],100,'Cantidades')
        self.assertEqual(len(rows),8)
        combinations={(r['zona'],r['género'],r['edad']) for r in rows}
        self.assertIn(('Sur','Mujer','30-44'),combinations)
        self.assertTrue(all(r['cantidad']==0 for r in rows))
        self.assertTrue(all(r['zona']=='Norte' for r in rows[:4]))
        reordered=server.make_quota_rows(frame,['género','zona','edad'],100,'Cantidades')
        self.assertTrue(all(r['género']=='Mujer' for r in reordered[:4]))

if __name__=='__main__': unittest.main()
